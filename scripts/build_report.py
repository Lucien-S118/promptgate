"""Build the editable submission report from the canonical Markdown text.

Optional authoring dependency: python-docx. This script is not needed for the app.
Render the resulting DOCX to PDF and inspect every page before distribution.
"""
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output/docx/PE6201_PromptGate_Tradeoff_Report_SIwen_Liu.docx'


def main():
    source = (ROOT / 'TRADEOFF_REPORT.md').read_text()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    section.footer_distance = Inches(.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.0
    doc.styles['Title'].font.size = Pt(20)
    doc.styles['Title'].paragraph_format.space_after = Pt(8)
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    doc.styles['Normal'].paragraph_format.widow_control = True
    metadata = []
    for block in source.strip().split('\n\n'):
        if block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('**Student:'):
            metadata = [re.sub(r'\*\*', '', line).strip() for line in block.splitlines()]
            p = doc.add_paragraph(' | '.join(metadata))
            p.paragraph_format.space_after = Pt(9)
            for run in p.runs:
                run.font.size = Pt(10)
        elif block.startswith('## '):
            doc.add_paragraph(block[3:], 'Heading 1')
        else:
            doc.add_paragraph(block.replace('\n', ' '))
    # Page numbers aid discussion of this multi-page report without adding branding.
    p = section.footer.paragraphs[0]
    p.alignment = 2
    p.add_run('Page ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    p._p.append(field)
    doc.core_properties.author = 'SIwen Liu'
    doc.core_properties.title = 'PromptGate business and technical tradeoffs'
    doc.core_properties.subject = 'PE6201 individual project'
    # The bundled base template can inherit a decorative Title paragraph border.
    # This academic report uses typography alone, so remove all paragraph borders.
    for tree in (doc.styles.element, doc.element):
        for border in tree.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    # Check conversion did not omit body paragraphs or sources.
    expected = []
    for block in source.strip().split('\n\n'):
        if block.startswith('**Student:'):
            expected.append(' | '.join(metadata))
        else:
            expected.append(re.sub(r'^#{1,2} ', '', block).replace('\n', ' '))
    actual = [p.text for p in Document(OUTPUT).paragraphs]
    if actual != expected:
        raise ValueError('Markdown-to-Word content mismatch')
    print(f'Created {OUTPUT}; {len(source.split())} whitespace-delimited words in source; all paragraphs verified.')


if __name__ == '__main__':
    main()
